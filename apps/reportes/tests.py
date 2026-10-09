from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.clientes.models import Cliente
from apps.gastos.models import RecuperacionGasto, RegistroFinanciero
from apps.honorarios.models import CobroHonorario, PagoHonorario
from apps.servicios.models import ClienteOcasional, PagoServicioTramite, ServicioTramite


class ReportesFinancierosTests(TestCase):
	def setUp(self):
		self.admin = get_user_model().objects.create_superuser(username='admin_reportes', password='ClaveLocal123!')
		self.client.force_login(self.admin)

	def test_resume_ingresos_egresos_reembolsos_y_flujo_neto(self):
		cliente = Cliente.objects.create(nombre='Cliente de reporte')
		for tipo_ingreso, monto, concepto in (
			(CobroHonorario.TipoIngreso.RECURRENTE, '100.00', 'Mensualidad'),
			(CobroHonorario.TipoIngreso.EXTRAORDINARIO, '40.00', 'Certificado'),
		):
			cobro = CobroHonorario.objects.create(
				cliente=cliente,
				periodo_tipo=CobroHonorario.Periodicidad.EXTRA if tipo_ingreso == CobroHonorario.TipoIngreso.EXTRAORDINARIO else CobroHonorario.Periodicidad.MENSUAL,
				tipo_ingreso=tipo_ingreso,
				anio=2026,
				periodo_numero=10,
				concepto=concepto,
				monto_total=Decimal(monto),
			)
			PagoHonorario.objects.create(cobro=cobro, monto=monto, fecha_pago=date(2026, 10, 8))

		cliente_ocasional = ClienteOcasional.objects.create(nombre='Cliente puntual')
		servicio = ServicioTramite.objects.create(
			cliente=cliente_ocasional,
			tipo=ServicioTramite.Tipo.LICENCIA,
			concepto='Licencia municipal',
			monto_total='30.00',
		)
		PagoServicioTramite.objects.create(servicio=servicio, monto='30.00', fecha_pago=date(2026, 10, 8))

		gasto = RegistroFinanciero.objects.create(
			concepto='Material de oficina',
			tipo=RegistroFinanciero.Tipo.GASTO,
			categoria=RegistroFinanciero.Categoria.SUMINISTROS,
			monto='80.00',
			fecha=date(2026, 10, 8),
			estado_pago=RegistroFinanciero.EstadoPago.PAGADO,
			recuperable=True,
			cliente=cliente,
		)
		RecuperacionGasto.objects.create(registro=gasto, monto='20.00', fecha=date(2026, 10, 9))
		RegistroFinanciero.objects.create(
			concepto='Servicio pendiente proveedor',
			tipo=RegistroFinanciero.Tipo.GASTO,
			categoria=RegistroFinanciero.Categoria.INTERNET,
			monto='50.00',
			fecha=date(2026, 10, 8),
			estado_pago=RegistroFinanciero.EstadoPago.PENDIENTE,
		)

		respuesta = self.client.get(reverse('reportes:index'), {'anio': '2026', 'mes': '10'})

		self.assertEqual(respuesta.status_code, 200)
		self.assertEqual(respuesta.context['ingresos_recurrentes'], Decimal('100.00'))
		self.assertEqual(respuesta.context['ingresos_extraordinarios'], Decimal('40.00'))
		self.assertEqual(respuesta.context['ingresos_tramites'], Decimal('30.00'))
		self.assertEqual(respuesta.context['recuperado_clientes'], Decimal('20.00'))
		self.assertEqual(respuesta.context['egresos_pagados'], Decimal('80.00'))
		self.assertEqual(respuesta.context['egresos_pendientes'], Decimal('50.00'))
		self.assertEqual(respuesta.context['flujo_neto'], Decimal('110.00'))
		self.assertContains(respuesta, 'Licencia municipal')
		self.assertContains(respuesta, 'css/clientes/clientes.css')
		self.assertContains(respuesta, 'class="button button--quiet"')
		self.assertContains(respuesta, 'class="button button--primary report-export-button"')

	def test_filtro_mensual_no_mezcla_movimientos_de_otro_mes(self):
		cliente = Cliente.objects.create(nombre='Cliente por mes')
		for mes, monto in ((9, '25.00'), (10, '60.00')):
			cobro = CobroHonorario.objects.create(
				cliente=cliente,
				periodo_tipo=CobroHonorario.Periodicidad.MENSUAL,
				anio=2026,
				periodo_numero=mes,
				concepto=f'Honorario {mes}',
				monto_total=Decimal(monto),
			)
			PagoHonorario.objects.create(cobro=cobro, monto=monto, fecha_pago=date(2026, mes, 8))

		respuesta = self.client.get(reverse('reportes:index'), {'anio': '2026', 'mes': '10'})

		self.assertEqual(respuesta.context['ingresos_recurrentes'], Decimal('60.00'))
		self.assertEqual(respuesta.context['total_movimientos'], 1)

	def test_excel_exporta_todas_las_hojas_y_respeta_el_mes(self):
		cliente = Cliente.objects.create(nombre='Cliente Excel')
		for mes, monto, concepto in ((9, '25.00', 'Pago fuera de período'), (10, '60.00', 'Pago del período')):
			cobro = CobroHonorario.objects.create(
				cliente=cliente,
				periodo_tipo=CobroHonorario.Periodicidad.MENSUAL,
				anio=2026,
				periodo_numero=mes,
				concepto=concepto,
				monto_total=Decimal(monto),
			)
			PagoHonorario.objects.create(cobro=cobro, monto=Decimal(monto), fecha_pago=date(2026, mes, 8))
		RegistroFinanciero.objects.create(
			concepto='Gasto Excel',
			tipo=RegistroFinanciero.Tipo.GASTO,
			categoria=RegistroFinanciero.Categoria.SUMINISTROS,
			monto='10.00',
			fecha=date(2026, 10, 8),
			estado_pago=RegistroFinanciero.EstadoPago.PAGADO,
		)
		RegistroFinanciero.objects.create(
			concepto='Pendiente proveedor',
			tipo=RegistroFinanciero.Tipo.GASTO,
			categoria=RegistroFinanciero.Categoria.INTERNET,
			monto='12.00',
			fecha=date(2026, 10, 9),
			estado_pago=RegistroFinanciero.EstadoPago.PENDIENTE,
		)

		respuesta = self.client.get(reverse('reportes:exportar_excel'), {'anio': '2026', 'mes': '10'})

		self.assertEqual(respuesta.status_code, 200)
		self.assertEqual(respuesta['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
		self.assertIn('reporte-financiero-2026-10.xlsx', respuesta['Content-Disposition'])
		libro = load_workbook(BytesIO(respuesta.content), data_only=True)
		self.assertEqual(libro.sheetnames, ['Resumen', 'Movimientos', 'Egresos', 'Flujo mensual'])
		for nombre_hoja, celda_interior in (
			('Resumen', 'A4'),
			('Movimientos', 'A2'),
			('Egresos', 'A2'),
			('Flujo mensual', 'A2'),
		):
			self.assertTrue(libro[nombre_hoja][celda_interior].fill.fgColor.rgb.endswith('EAF5EE'))
		resumen = {fila[0]: fila[1] for fila in libro['Resumen'].iter_rows(min_row=4, values_only=True)}
		self.assertEqual(resumen['Honorarios recurrentes cobrados'], 60)
		self.assertEqual(resumen['Egresos pagados'], 10)
		movimientos = [fila[3] for fila in libro['Movimientos'].iter_rows(min_row=2, values_only=True)]
		self.assertIn('Pago del período', movimientos)
		self.assertNotIn('Pago fuera de período', movimientos)
		egresos = list(libro['Egresos'].iter_rows(min_row=2, values_only=True))
		self.assertTrue(any(fila[2] == 'Pendiente proveedor' and fila[5] == 'Pendiente de pago' for fila in egresos))
