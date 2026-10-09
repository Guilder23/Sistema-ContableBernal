from decimal import Decimal
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.gastos.models import RecuperacionGasto, RegistroFinanciero
from apps.honorarios.models import CobroHonorario, PagoHonorario
from apps.servicios.models import PagoServicioTramite, ServicioTramite


MESES = (
    (1, 'Enero'), (2, 'Febrero'), (3, 'Marzo'), (4, 'Abril'),
    (5, 'Mayo'), (6, 'Junio'), (7, 'Julio'), (8, 'Agosto'),
    (9, 'Septiembre'), (10, 'Octubre'), (11, 'Noviembre'), (12, 'Diciembre'),
)
CERO = Decimal('0.00')


def _total(queryset):
    return queryset.aggregate(total=Sum('monto'))['total'] or CERO


def _totales_por_mes(queryset, campo_fecha):
    campo_mes = f'{campo_fecha}__month'
    return {
        item[campo_mes]: item['total']
        for item in queryset.values(campo_mes).annotate(total=Sum('monto'))
    }


def _consultas_periodo(anio, mes_numero=None):
    pagos_honorarios_base = PagoHonorario.objects.select_related(
        'cobro', 'cobro__cliente', 'registrado_por',
    ).filter(fecha_pago__year=anio)
    pagos_servicios_base = PagoServicioTramite.objects.select_related(
        'servicio', 'servicio__cliente', 'registrado_por',
    ).filter(fecha_pago__year=anio)
    recuperaciones_base = RecuperacionGasto.objects.select_related(
        'registro', 'registro__cliente', 'registrado_por',
    ).filter(fecha__year=anio)
    registros_base = RegistroFinanciero.objects.select_related('cliente', 'creado_por').filter(fecha__year=anio)

    if mes_numero:
        pagos_honorarios = pagos_honorarios_base.filter(fecha_pago__month=mes_numero)
        pagos_servicios = pagos_servicios_base.filter(fecha_pago__month=mes_numero)
        recuperaciones = recuperaciones_base.filter(fecha__month=mes_numero)
        registros = registros_base.filter(fecha__month=mes_numero)
    else:
        pagos_honorarios = pagos_honorarios_base
        pagos_servicios = pagos_servicios_base
        recuperaciones = recuperaciones_base
        registros = registros_base

    return (
        pagos_honorarios_base, pagos_servicios_base, recuperaciones_base, registros_base,
        pagos_honorarios, pagos_servicios, recuperaciones, registros,
    )


def _movimientos_periodo(pagos_honorarios, pagos_servicios, recuperaciones, registros, limite_por_fuente=None):
    movimientos = []
    for pago in pagos_honorarios.order_by('-fecha_pago', '-pk')[:limite_por_fuente]:
        movimientos.append({
            'fecha': pago.fecha_pago,
            'clase': 'entrada',
            'origen': pago.cobro.get_tipo_ingreso_display(),
            'concepto': pago.cobro.concepto,
            'persona': pago.cobro.cliente.nombre,
            'monto': pago.monto,
            'estado': 'Recibido',
        })
    for pago in pagos_servicios.order_by('-fecha_pago', '-pk')[:limite_por_fuente]:
        movimientos.append({
            'fecha': pago.fecha_pago,
            'clase': 'entrada',
            'origen': 'Trámite ocasional',
            'concepto': pago.servicio.concepto,
            'persona': pago.servicio.cliente.nombre,
            'monto': pago.monto,
            'estado': 'Recibido',
        })
    for recuperacion in recuperaciones.order_by('-fecha', '-pk')[:limite_por_fuente]:
        movimientos.append({
            'fecha': recuperacion.fecha,
            'clase': 'recuperacion',
            'origen': 'Reembolso de cliente',
            'concepto': recuperacion.registro.concepto,
            'persona': recuperacion.registro.cliente.nombre if recuperacion.registro.cliente else 'Cliente',
            'monto': recuperacion.monto,
            'estado': 'Recuperado',
        })
    for registro in registros.order_by('-fecha', '-pk')[:limite_por_fuente]:
        movimientos.append({
            'fecha': registro.fecha,
            'clase': 'salida',
            'origen': registro.get_tipo_display(),
            'concepto': registro.concepto,
            'persona': registro.proveedor or (registro.cliente.nombre if registro.cliente else 'Estudio'),
            'monto': registro.monto,
            'estado': registro.get_estado_pago_display(),
        })
    movimientos.sort(key=lambda item: item['fecha'], reverse=True)
    return movimientos


def _porcentaje(valor, maximo):
    return min(100, int(valor / maximo * 100)) if maximo else 0


def _periodo_desde_request(request):
    hoy = timezone.localdate()
    try:
        anio = int(request.GET.get('anio', hoy.year))
        if not 2000 <= anio <= 2200:
            raise ValueError
    except (TypeError, ValueError):
        anio = hoy.year
    mes = request.GET.get('mes', '').strip()
    if not mes.isdigit() or not 1 <= int(mes) <= 12:
        return anio, '', None
    return anio, mes, int(mes)


def _texto_excel(valor):
    texto = str(valor or '')
    return f"'{texto}" if texto.startswith(('=', '+', '-', '@')) else texto


def _formatear_hoja(hoja, anchos, fila_encabezado=1):
    hoja.freeze_panes = f'A{fila_encabezado + 1}'
    hoja.auto_filter.ref = hoja.dimensions
    relleno_interior = PatternFill(fill_type='solid', fgColor='EAF5EE')
    for celda in hoja[fila_encabezado]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill(fill_type='solid', fgColor='355E4B')
        celda.alignment = Alignment(vertical='center')
    for fila in hoja.iter_rows(min_row=fila_encabezado + 1):
        for celda in fila:
            celda.fill = relleno_interior
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho


@login_required
@require_GET
def exportar_excel(request):
    anio, mes, mes_numero = _periodo_desde_request(request)
    (
        _, _, _, _,
        pagos_honorarios, pagos_servicios, recuperaciones, registros,
    ) = _consultas_periodo(anio, mes_numero)
    movimientos = _movimientos_periodo(pagos_honorarios, pagos_servicios, recuperaciones, registros)
    recurrentes = _total(pagos_honorarios.filter(cobro__tipo_ingreso=CobroHonorario.TipoIngreso.RECURRENTE))
    extraordinarios = _total(pagos_honorarios.filter(cobro__tipo_ingreso=CobroHonorario.TipoIngreso.EXTRAORDINARIO))
    tramites = _total(pagos_servicios)
    recuperado = _total(recuperaciones)
    pagados = registros.filter(estado_pago=RegistroFinanciero.EstadoPago.PAGADO)
    pendientes = registros.filter(estado_pago=RegistroFinanciero.EstadoPago.PENDIENTE)
    total_ingresos = recurrentes + extraordinarios + tramites
    total_egresos = _total(pagados)

    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    resumen.merge_cells('A1:B1')
    resumen['A1'] = 'Reporte financiero'
    resumen['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    resumen['A1'].fill = PatternFill(fill_type='solid', fgColor='355E4B')
    resumen['A1'].alignment = Alignment(vertical='center')
    resumen.row_dimensions[1].height = 26
    periodo = f'{dict(MESES).get(mes_numero)} {anio}' if mes_numero else f'Año {anio}'
    resumen.append([])
    resumen.append(['Concepto', 'Monto (Bs)'])
    resumen.append(['Período', periodo])
    resumen.append(['Honorarios recurrentes cobrados', recurrentes])
    resumen.append(['Servicios extraordinarios cobrados', extraordinarios])
    resumen.append(['Trámites ocasionales cobrados', tramites])
    resumen.append(['Total ingresos cobrados', total_ingresos])
    resumen.append(['Reembolsos recibidos de clientes', recuperado])
    resumen.append(['Egresos pagados', total_egresos])
    resumen.append(['Flujo neto de caja', total_ingresos + recuperado - total_egresos])
    resumen.append(['Egresos pendientes de pago', _total(pendientes)])
    _formatear_hoja(resumen, {'A': 42, 'B': 22}, fila_encabezado=3)
    for fila in range(4, resumen.max_row + 1):
        if isinstance(resumen.cell(fila, 2).value, (int, float, Decimal)):
            resumen.cell(fila, 2).number_format = '"Bs" #,##0.00'

    hoja_movimientos = libro.create_sheet('Movimientos')
    hoja_movimientos.append(['Fecha', 'Movimiento', 'Origen', 'Concepto', 'Cliente / proveedor', 'Estado', 'Monto (Bs)'])
    etiquetas = {'entrada': 'Ingreso', 'recuperacion': 'Reembolso', 'salida': 'Egreso'}
    for item in movimientos:
        hoja_movimientos.append([
            item['fecha'],
            etiquetas[item['clase']],
            _texto_excel(item['origen']),
            _texto_excel(item['concepto']),
            _texto_excel(item['persona']),
            item['estado'],
            float(item['monto']),
        ])
    _formatear_hoja(hoja_movimientos, {'A': 14, 'B': 16, 'C': 24, 'D': 38, 'E': 30, 'F': 20, 'G': 18})
    for fila in range(2, hoja_movimientos.max_row + 1):
        hoja_movimientos.cell(fila, 1).number_format = 'dd/mm/yyyy'
        hoja_movimientos.cell(fila, 7).number_format = '"Bs" #,##0.00'

    egresos_categoria = libro.create_sheet('Egresos')
    egresos_categoria.append(['Tipo', 'Categoría', 'Concepto', 'Proveedor / cliente', 'Fecha', 'Estado', 'Monto (Bs)'])
    for registro in registros.order_by('-fecha', '-pk'):
        egresos_categoria.append([
            registro.get_tipo_display(),
            registro.get_categoria_display(),
            _texto_excel(registro.concepto),
            _texto_excel(registro.proveedor or (registro.cliente.nombre if registro.cliente else 'Estudio')),
            registro.fecha,
            registro.get_estado_pago_display(),
            float(registro.monto),
        ])
    _formatear_hoja(egresos_categoria, {'A': 18, 'B': 26, 'C': 38, 'D': 30, 'E': 14, 'F': 22, 'G': 18})
    for fila in range(2, egresos_categoria.max_row + 1):
        egresos_categoria.cell(fila, 5).number_format = 'dd/mm/yyyy'
        egresos_categoria.cell(fila, 7).number_format = '"Bs" #,##0.00'

    meses_flujo = libro.create_sheet('Flujo mensual')
    meses_flujo.append(['Mes', 'Ingresos y reembolsos (Bs)', 'Egresos pagados (Bs)', 'Flujo neto (Bs)'])
    meses = ((mes_numero, dict(MESES)[mes_numero]),) if mes_numero else MESES
    (
        pagos_honorarios_base, pagos_servicios_base, recuperaciones_base, registros_base,
        _, _, _, _,
    ) = _consultas_periodo(anio)
    totales_honorarios = _totales_por_mes(pagos_honorarios_base, 'fecha_pago')
    totales_tramites = _totales_por_mes(pagos_servicios_base, 'fecha_pago')
    totales_recuperaciones = _totales_por_mes(recuperaciones_base, 'fecha')
    totales_egresos = _totales_por_mes(registros_base.filter(estado_pago=RegistroFinanciero.EstadoPago.PAGADO), 'fecha')
    for numero, nombre in meses:
        entradas = totales_honorarios.get(numero, CERO) + totales_tramites.get(numero, CERO) + totales_recuperaciones.get(numero, CERO)
        salidas = totales_egresos.get(numero, CERO)
        meses_flujo.append([nombre, float(entradas), float(salidas), float(entradas - salidas)])
    _formatear_hoja(meses_flujo, {'A': 18, 'B': 30, 'C': 26, 'D': 22})
    for fila in range(2, meses_flujo.max_row + 1):
        for columna in range(2, 5):
            meses_flujo.cell(fila, columna).number_format = '"Bs" #,##0.00'

    salida = BytesIO()
    libro.save(salida)
    nombre_periodo = f'{anio}-{mes_numero:02d}' if mes_numero else str(anio)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="reporte-financiero-{nombre_periodo}.xlsx"'
    return respuesta


@login_required
def index(request):
    hoy = timezone.localdate()
    anio, mes, mes_numero = _periodo_desde_request(request)

    (
        pagos_honorarios_base, pagos_servicios_base, recuperaciones_base, registros_base,
        pagos_honorarios, pagos_servicios, recuperaciones, registros,
    ) = _consultas_periodo(anio, mes_numero)

    pagos_recurrentes = pagos_honorarios.filter(cobro__tipo_ingreso=CobroHonorario.TipoIngreso.RECURRENTE)
    pagos_extraordinarios = pagos_honorarios.filter(cobro__tipo_ingreso=CobroHonorario.TipoIngreso.EXTRAORDINARIO)
    ingresos_recurrentes = _total(pagos_recurrentes)
    ingresos_extraordinarios = _total(pagos_extraordinarios)
    ingresos_tramites = _total(pagos_servicios)
    recuperado_clientes = _total(recuperaciones)
    ingresos_servicios = ingresos_recurrentes + ingresos_extraordinarios + ingresos_tramites
    egresos_pagados = registros.filter(estado_pago=RegistroFinanciero.EstadoPago.PAGADO)
    egresos_pendientes = registros.filter(estado_pago=RegistroFinanciero.EstadoPago.PENDIENTE)
    total_egresos_pagados = _total(egresos_pagados)
    total_egresos_pendientes = _total(egresos_pendientes)
    flujo_neto = ingresos_servicios + recuperado_clientes - total_egresos_pagados

    tipos_labels = dict(RegistroFinanciero.Tipo.choices)
    eg_resumen = list(
        egresos_pagados.values('tipo').annotate(total=Sum('monto')).order_by('-total')
    )
    eg_resumen = [
        {'nombre': tipos_labels.get(item['tipo'], item['tipo']), 'monto': item['total']}
        for item in eg_resumen
    ]
    maximo_tipo = max((item['monto'] for item in eg_resumen), default=CERO)
    for item in eg_resumen:
        item['porcentaje'] = _porcentaje(item['monto'], maximo_tipo)

    categorias_labels = dict(RegistroFinanciero.Categoria.choices)
    eg_categorias = list(
        egresos_pagados.values('categoria').annotate(total=Sum('monto')).order_by('-total')[:8]
    )
    eg_categorias = [
        {'nombre': categorias_labels.get(item['categoria'], item['categoria']), 'monto': item['total']}
        for item in eg_categorias
    ]
    maximo_categoria = max((item['monto'] for item in eg_categorias), default=CERO)
    for item in eg_categorias:
        item['porcentaje'] = _porcentaje(item['monto'], maximo_categoria)

    meses_reporte = []
    maximo_entradas = CERO
    maximo_salidas = CERO
    honorarios_por_mes = _totales_por_mes(pagos_honorarios_base, 'fecha_pago')
    tramites_por_mes = _totales_por_mes(pagos_servicios_base, 'fecha_pago')
    recuperaciones_por_mes = _totales_por_mes(recuperaciones_base, 'fecha')
    egresos_por_mes = _totales_por_mes(
        registros_base.filter(estado_pago=RegistroFinanciero.EstadoPago.PAGADO),
        'fecha',
    )
    for numero, nombre in MESES:
        entradas = (
            honorarios_por_mes.get(numero, CERO)
            + tramites_por_mes.get(numero, CERO)
            + recuperaciones_por_mes.get(numero, CERO)
        )
        salidas = egresos_por_mes.get(numero, CERO)
        meses_reporte.append({'nombre': nombre, 'entradas': entradas, 'salidas': salidas})
        maximo_entradas = max(maximo_entradas, entradas)
        maximo_salidas = max(maximo_salidas, salidas)
    for item in meses_reporte:
        item['entradas_pct'] = _porcentaje(item['entradas'], maximo_entradas)
        item['salidas_pct'] = _porcentaje(item['salidas'], maximo_salidas)

    clientes_por_cobrar = CobroHonorario.objects.exclude(estado=CobroHonorario.Estado.CANCELADO).prefetch_related('pagos')
    servicios_por_cobrar = ServicioTramite.objects.prefetch_related('pagos')
    saldo_honorarios = sum((cobro.saldo_pendiente for cobro in clientes_por_cobrar), CERO)
    saldo_tramites = sum((servicio.saldo_pendiente for servicio in servicios_por_cobrar), CERO)
    registros_recuperables = RegistroFinanciero.objects.filter(recuperable=True).prefetch_related('recuperaciones')
    saldo_recuperar = sum((registro.saldo_por_recuperar for registro in registros_recuperables), CERO)

    movimientos = _movimientos_periodo(
        pagos_honorarios,
        pagos_servicios,
        recuperaciones,
        registros,
        limite_por_fuente=100,
    )[:100]
    total_movimientos = pagos_honorarios.count() + pagos_servicios.count() + recuperaciones.count() + registros.count()

    query_params = request.GET.copy()
    return render(request, 'reportes/reportes.html', {
        'titulo_modulo': 'Reportes',
        'descripcion_modulo': 'Ingresos cobrados, egresos, recuperaciones y flujo de caja.',
        'modulo_activo': 'reportes',
        'hoy': hoy,
        'anio': anio,
        'mes': mes,
        'meses': MESES,
        'ingresos_recurrentes': ingresos_recurrentes,
        'ingresos_extraordinarios': ingresos_extraordinarios,
        'ingresos_tramites': ingresos_tramites,
        'recuperado_clientes': recuperado_clientes,
        'ingresos_servicios': ingresos_servicios,
        'egresos_pagados': total_egresos_pagados,
        'egresos_pendientes': total_egresos_pendientes,
        'flujo_neto': flujo_neto,
        'eg_resumen': eg_resumen,
        'eg_categorias': eg_categorias,
        'meses_reporte': meses_reporte,
        'saldo_honorarios': saldo_honorarios,
        'saldo_tramites': saldo_tramites,
        'saldo_recuperar': saldo_recuperar,
        'movimientos': movimientos[:100],
        'total_movimientos': total_movimientos,
        'query_string': query_params.urlencode(),
    })
